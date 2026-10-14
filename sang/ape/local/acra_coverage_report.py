"""Coverage of Python rules against JTBD jobs for the Acra DSCR guidelines (local, a report).

Author: Virendra Mehta <virendra.mehta@jazzx.ai>

Reads the four extraction sets from the local Plato (tenant acra-lending, pack acra-dscr-policy):
JTBD Standard, JTBD Platinum, Python Standard, Python Platinum. For each guideline document it
aligns every Python rule to the JTBD jobs it implements (lexical candidates, then a model tie-break
per rule), classifies each job (covered / computable-but-missing / not-computable), lists the
Python rules no job accounts for, and writes JSON and xlsx with aggregates and the model cost.

    python scripts/local/acra_coverage_report.py [--out /tmp/cov] [--budget 5.0]

Guideline documents carry no borrower data. The model is the deployment's default.
"""

from __future__ import annotations

import argparse
import asyncio
import json
import math
import os
import re
import sys
from collections import Counter
from pathlib import Path

import httpx

ROOT = Path("/Users/sangit/src/japes")
sys.path.insert(0, str(ROOT))

API = "http://localhost:8000/api/v1"
TENANT = "acra-lending"
PACK = os.environ.get("ACRA_PACK", "acra-dscr-policy")
SCHEMA = ROOT / "plato" / "data" / "seed_packs" / "acra_dscr_policy" / "schemas" / "loan_application.json"
CANDIDATES = 5
EXCERPT_CHARS = 1800
#: Lines of a rule's code the model sees: all of every generated rule (the longest was 173).
CODE_LINES = 400
STOP = set("the and for that with this are not any must may has have from per shall will than then "
           "all its into over each been being their there which when where who whose does done "
           "only also such these those other more less must".split())


ACRA_JTBD = {"standard": ROOT / "documents" / "Acra" / "acra_New_Guideline_JTBD.json",
             "platinum": ROOT / "documents" / "Acra" / "acra_New_Guideline_JTBD_Platinum.json"}


def _fetch(run_ids: set[str] | None = None) -> list[dict]:
    """The pack's proposals; with ``run_ids``, only those a listed run last proposed."""
    with httpx.Client(timeout=30, headers={"x-tenant-id": TENANT}) as http:
        proposals = http.get(f"{API}/authoring/review/{PACK}/proposals").json()
    return [p for p in proposals if run_ids is None or p.get("run_id") in run_ids]


def _final_run_ids(path: Path) -> set[str]:
    """Every succeeded attempt in an `acra_rerun_extractions.py` output. A proposal carries the run
    that last proposed it, and names are per segment, so a part one attempt lost and another
    extracted is counted once."""
    return {a["run_id"] for job in json.loads(path.read_text()) for a in job["attempts"]
            if a["status"] == "succeeded"}


def _still_failed(job: dict) -> list[str]:
    """The parts every succeeded attempt of ``job`` failed."""
    done = [set(a["failed_segments"]) for a in job["attempts"] if a["status"] == "succeeded"]
    return sorted(set.intersection(*done)) if done else ["(no attempt succeeded)"]


def _acra_jobs(doc: str) -> list[dict]:
    """Acra's JTBD file for ``doc``, each job shaped like a procedure proposal: the purpose is the
    skill, the step the COMMON requirement text, the reads the metric keys."""
    artifact = json.loads(ACRA_JTBD[doc].read_text())
    jobs = []
    for value in artifact["values"]:
        text = next((r.get("requirement_text", "") for r in value.get("requirements") or []
                     if r.get("investor") == "COMMON"), "")
        reads = [m["metric_key"] for m in value.get("loan_metrics") or [] if m.get("metric_key")]
        jobs.append({"name": f"acra-{doc}.jtbd.{value['mnemonic']}", "excerpt": text,
                     "reasoning": f"from section '{value.get('section', '')}'",
                     "rule": {"condition": {"kind": "procedure", "purpose": value.get("skill", ""),
                                            "steps": [text], "checks": [], "reads": reads}}})
    return jobs


def _kind(p: dict) -> str:
    cond = p["rule"]["condition"]
    return cond.get("kind", "?") if isinstance(cond, dict) else "?"


def _doc(p: dict) -> str:
    return "standard" if p["name"].startswith("acra-standard.") else "platinum"


def _tokens(text: str) -> list[str]:
    return [t for t in re.findall(r"[a-z][a-z0-9_]{2,}", text.lower()) if t not in STOP]


def _numbers(text: str) -> set[str]:
    return {n.replace(",", "").replace("_", "").rstrip(".") for n in re.findall(r"\d[\d,_]*\.?\d*", text)
            if len(n.replace(",", "").replace("_", "")) >= 2}


def _job_text(p: dict) -> str:
    c = p["rule"]["condition"]
    return " ".join([p["name"].rsplit(".", 1)[-1].replace("_", " "), c.get("purpose", ""),
                     " ".join(c.get("steps") or []), " ".join(c.get("checks") or []),
                     p.get("excerpt", "")])


def _py_text(p: dict) -> str:
    c = p["rule"]["condition"]
    return " ".join([p.get("excerpt", ""), p.get("reasoning", ""), c.get("source", "")])


def _tfidf(docs: list[list[str]]) -> tuple[list[dict[str, float]], dict[str, float]]:
    df = Counter(t for d in docs for t in set(d))
    idf = {t: math.log((1 + len(docs)) / (1 + n)) + 1 for t, n in df.items()}
    vecs = []
    for d in docs:
        tf = Counter(d)
        v = {t: (1 + math.log(c)) * idf[t] for t, c in tf.items()}
        norm = math.sqrt(sum(x * x for x in v.values())) or 1.0
        vecs.append({t: x / norm for t, x in v.items()})
    return vecs, idf


def _cos(a: dict[str, float], b: dict[str, float]) -> float:
    return sum(x * b.get(t, 0.0) for t, x in a.items())


def candidates(jobs: list[dict], rules: list[dict]) -> dict[int, list[tuple[int, float]]]:
    """For each rule index, its best jobs: tf-idf over the clause text, reads and numbers."""
    docs = [_tokens(_job_text(j)) for j in jobs] + [_tokens(_py_text(r)) for r in rules]
    vecs, _ = _tfidf(docs)
    jv, rv = vecs[:len(jobs)], vecs[len(jobs):]
    out = {}
    for i, rule in enumerate(rules):
        reads = set(rule["rule"]["condition"].get("reads") or [])
        nums = _numbers(_py_text(rule))
        scored = []
        for k, job in enumerate(jobs):
            jr = set(job["rule"]["condition"].get("reads") or [])
            jacc = len(reads & jr) / len(reads | jr) if reads | jr else 0.0
            jn = _numbers(_job_text(job))
            numo = len(nums & jn) / len(nums | jn) if nums | jn else 0.0
            scored.append((k, 0.6 * _cos(rv[i], jv[k]) + 0.3 * jacc + 0.1 * numo))
        out[i] = sorted((s for s in scored if s[1] > 0), key=lambda s: -s[1])[:CANDIDATES]
    return out


class Tally:
    def __init__(self, service, budget: float):
        self._service, self.budget = service, budget
        self.calls = self.input = self.output = 0
        self.cost = 0.0
        self.upper = False
        self.stopped = False

    async def run(self, **kw):
        if self.cost >= self.budget:
            self.stopped = True
            raise RuntimeError("budget spent")
        result = await self._service.run(**kw)
        trace = getattr(result, "trace", None)
        self.calls += 1
        if trace is not None:
            self.input += trace.tokens_input
            self.output += trace.tokens_output
            self.cost += trace.cost_usd
            self.upper |= bool(getattr(trace, "cost_is_upper_bound", False))
        return result

    def summary(self) -> dict:
        return {"calls": self.calls, "input_tokens": self.input, "output_tokens": self.output,
                "cost_usd": round(self.cost, 4), "cost_is_upper_bound": self.upper,
                "budget_usd": self.budget, "stopped_on_budget": self.stopped}


def _schemas():
    from pydantic import BaseModel, Field

    class Match(BaseModel):
        job_ids: list[int] = Field(default_factory=list)
        relation: str = "no_job"     # implements | partial | duplicate | no_job
        note: str = ""

    class Computable(BaseModel):
        computable: bool = False
        missing: list[str] = Field(default_factory=list)
        why: str = ""

    return Match, Computable


MATCH_SYSTEM = (
    "You align a generated Python rule to the JTBD jobs of the same underwriting guideline. A job is a "
    "checkable requirement. Choose the job numbers the Python rule implements: all of a job (implements) "
    "or part of one (partial). Return relation 'duplicate' when the rule repeats another listed job's "
    "check under another name, and 'no_job' with no job_ids when no listed job states what the rule "
    "checks (the rule may invent a threshold the text does not print). Use only the numbers shown.")
COMPUTE_SYSTEM = (
    "A JTBD job names a guideline requirement and the loan fields it reads. Say whether a deterministic "
    "rule over exactly those fields can decide it: computable=true when the checks reduce to comparisons "
    "over the listed fields; false when it needs judgment, a document's content, an external lookup, or "
    "a field that is not listed (name those in `missing`). One short sentence in `why`.")
COMPUTE_SYSTEM_FOREIGN = (
    "A JTBD job names a guideline requirement and the metric keys it reads, in another vocabulary than "
    "the loan application schema. Say whether a deterministic rule over schema fields can decide it: "
    "computable=true when the checks reduce to comparisons over fields the schema has (under any name); "
    "false when it needs judgment, a document's content, an external lookup, or a field the schema "
    "lacks (name those in `missing`). One short sentence in `why`.")


def _job_card(i: int, job: dict) -> str:
    c = job["rule"]["condition"]
    return (f"[{i}] {job['name'].rsplit('.', 1)[-1].replace('_', ' ')}: {c.get('purpose', '')} "
            f"Checks: {' | '.join((c.get('checks') or [])[:4])} Reads: {', '.join(c.get('reads') or [])}")


async def tie_break(agent, rule: dict, cands: list[tuple[int, float]], jobs: list[dict], Match):
    code = "\n".join(rule["rule"]["condition"].get("source", "").splitlines()[:CODE_LINES])
    prompt = (f"PYTHON RULE\nClause text:\n{rule.get('excerpt', '')[:EXCERPT_CHARS]}\n\nReads: "
              f"{', '.join(rule['rule']['condition'].get('reads') or [])}\nCode:\n{code}\n\n"
              f"CANDIDATE JOBS\n" + "\n".join(_job_card(k, jobs[k]) for k, _ in cands))
    try:
        result = await agent.run(system_prompt=MATCH_SYSTEM,
                                 messages=[{"role": "user", "content": prompt}],
                                 output_schema=Match, temperature=0.0)
    except Exception as exc:  # noqa: BLE001
        return Match(relation="unknown", note=f"{type(exc).__name__}"), False
    out = result.output if isinstance(result.output, Match) else Match(**result.output)
    out.job_ids = [i for i in out.job_ids if i in {k for k, _ in cands}]
    return out, True


#: Set for Acra's jobs, whose metric keys name fields in Acra's vocabulary, not the schema's.
FOREIGN_VOCABULARY = False


async def computability(agent, job: dict, schema_fields: set[str], Computable):
    c = job["rule"]["condition"]
    reads = c.get("reads") or []
    if FOREIGN_VOCABULARY:
        reading = (f"The job's metric keys, in another vocabulary (map each to a schema field by "
                   f"meaning, e.g. LoanCore.cltv_percent is cltv_pct; the job may also need schema "
                   f"fields its keys do not name): {', '.join(reads) or '(none)'}")
    else:
        reading = f"The job reads: {', '.join(r for r in reads if r in schema_fields) or '(none)'}"
    prompt = (f"JOB: {_job_card(0, job)}\nPurpose: {c.get('purpose', '')}\nSteps: "
              f"{' | '.join(c.get('steps') or [])}\nFields the loan application schema has: "
              f"{', '.join(sorted(schema_fields))}\n{reading}")
    try:
        result = await agent.run(system_prompt=COMPUTE_SYSTEM_FOREIGN if FOREIGN_VOCABULARY else COMPUTE_SYSTEM,
                                 messages=[{"role": "user", "content": prompt}],
                                 output_schema=Computable, temperature=0.0)
    except Exception as exc:  # noqa: BLE001
        return Computable(why=f"unclassified ({type(exc).__name__})"), False
    return (result.output if isinstance(result.output, Computable) else Computable(**result.output)), True


async def analyse(doc: str, jobs: list[dict], rules: list[dict], agent, schema_fields, Match, Computable):
    cand = candidates(jobs, rules)
    aligned: dict[int, list[dict]] = {}
    rule_rows = []
    for i, rule in enumerate(rules):
        match, ok = await tie_break(agent, rule, cand[i], jobs, Match)
        for k in match.job_ids:
            aligned.setdefault(k, []).append({"rule": i, "relation": match.relation})
        code_nums = _numbers(rule["rule"]["condition"].get("source", ""))
        text_nums = _numbers(rule.get("excerpt", ""))
        rule_rows.append({
            "rule": rule["name"].rsplit(".", 1)[-1], "citation": (rule["rule"]["citations"] or [""])[0],
            "reads": rule["rule"]["condition"].get("reads") or [], "relation": match.relation,
            "jobs": [jobs[k]["name"].rsplit(".", 1)[-1] for k in match.job_ids], "note": match.note,
            "numbers_not_in_clause": sorted(code_nums - text_nums)[:12], "model_decided": ok,
            "top_lexical": [(jobs[k]["name"].rsplit(".", 1)[-1], round(s, 3)) for k, s in cand[i][:2]]})
    job_rows = []
    for k, job in enumerate(jobs):
        c = job["rule"]["condition"]
        reads = c.get("reads") or []
        row = {"job": job["name"].rsplit(".", 1)[-1], "section": re.sub(
            r".*from section '(.*)'$", r"\1", job.get("reasoning", "")), "reads": reads}
        if k in aligned:
            full = any(a["relation"] == "implements" for a in aligned[k])
            row.update(status="covered" if full else "covered_partial",
                       python_rules=[rule_rows[a["rule"]]["rule"] for a in aligned[k]])
        elif not reads and not FOREIGN_VOCABULARY:
            row.update(status="not_computable", why="reads no loan field")
        else:
            comp, ok = await computability(agent, job, schema_fields, Computable)
            row.update(status="computable_but_missing" if comp.computable else "not_computable",
                       why=comp.why, missing=comp.missing, model_decided=ok)
        job_rows.append(row)
    orphans = [r for r in rule_rows if not r["jobs"]]
    status = Counter(r["status"] for r in job_rows)
    return {"document": doc, "jobs": len(jobs), "python_rules": len(rules),
            "job_status": dict(status), "orphan_python_rules": len(orphans),
            "python_rules_aligned": len(rules) - len(orphans),
            "covered_share": round((status["covered"] + status["covered_partial"]) / len(jobs), 3) if jobs else 0,
            "job_rows": job_rows, "rule_rows": rule_rows}


def _xlsx(report: dict, path: Path) -> None:
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = "Summary"
    ws.append(["document", "jobs", "python rules", "covered", "covered_partial",
               "computable_but_missing", "not_computable", "covered share", "orphan python rules"])
    for d in report["documents"]:
        s = d["job_status"]
        ws.append([d["document"], d["jobs"], d["python_rules"], s.get("covered", 0),
                   s.get("covered_partial", 0), s.get("computable_but_missing", 0),
                   s.get("not_computable", 0), d["covered_share"], d["orphan_python_rules"]])
    ws.append([])
    ws.append(["model cost"] + [f"{k}: {v}" for k, v in report["cost"].items()])
    for d in report["documents"]:
        sheet = wb.create_sheet(f"{d['document']} jobs")
        sheet.append(["job", "section", "status", "python rules", "reads", "why / missing"])
        for r in d["job_rows"]:
            sheet.append([r["job"], r["section"], r["status"], ", ".join(r.get("python_rules", [])),
                          ", ".join(r["reads"]), (r.get("why", "") + " " + ", ".join(r.get("missing", []))).strip()])
        sheet = wb.create_sheet(f"{d['document']} python")
        sheet.append(["rule", "citation", "relation", "jobs", "reads", "numbers not in clause", "note"])
        for r in d["rule_rows"]:
            sheet.append([r["rule"], r["citation"], r["relation"], ", ".join(r["jobs"]),
                          ", ".join(r["reads"]), ", ".join(r["numbers_not_in_clause"]), r["note"]])
    sheet = wb.create_sheet("Failed segments")
    sheet.append(["run", "representation", "segment", "reason"])
    for f in report["failed_segments"]:
        sheet.append([f["run"], f["representation"], f["segment"], f.get("reason", "")])
    wb.save(path)


async def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--out", default="/tmp/cov")
    parser.add_argument("--budget", type=float, default=5.0)
    parser.add_argument("--runs", default="", help="an acra_rerun_extractions.py output: only its runs")
    parser.add_argument("--jobs", choices=("ours", "acra"), default="ours",
                        help="the job side: our extracted JTBDs, or Acra's JTBD files")
    args = parser.parse_args()
    out = Path(args.out)
    out.mkdir(parents=True, exist_ok=True)

    from jazzx_sdk.agents.service import AgentExecutionService

    run_ids = _final_run_ids(Path(args.runs)) if args.runs else None
    proposals = [p for p in _fetch(run_ids) if p["kind"] == "rule"]
    sets = {(doc, kind): [p for p in proposals if _doc(p) == doc and _kind(p) == kind]
            for doc in ("standard", "platinum") for kind in ("procedure", "python")}
    if args.jobs == "acra":
        global FOREIGN_VOCABULARY
        FOREIGN_VOCABULARY = True
        sets.update({(doc, "procedure"): _acra_jobs(doc) for doc in ("standard", "platinum")})
    schema_fields = set(json.loads(SCHEMA.read_text())["properties"])
    Match, Computable = _schemas()
    agent = Tally(AgentExecutionService(), args.budget)
    documents = []
    for doc in ("standard", "platinum"):
        jobs, rules = sets[(doc, "procedure")], sets[(doc, "python")]
        print(f"{doc}: {len(jobs)} jobs, {len(rules)} python rules", flush=True)
        documents.append(await analyse(doc, jobs, rules, agent, schema_fields, Match, Computable))
        print(f"  cost so far ${agent.cost:.4f}", flush=True)
    if args.runs:
        failed = [{"run": " ".join([*(a["run_id"][:8] for a in job["attempts"]),
                                    job["representation"], job["document"]]),
                   "representation": job["representation"], "segment": segment}
                  for job in json.loads(Path(args.runs).read_text())
                  for segment in _still_failed(job)]
    else:
        failed = json.loads((out / "failed_segments.json").read_text()) if (out / "failed_segments.json").exists() else []
    report = {"tenant": TENANT, "pack": PACK, "jobs_from": args.jobs, "runs": sorted(run_ids or []),
              "documents": documents, "cost": agent.summary(), "failed_segments": failed}
    (out / "coverage_report.json").write_text(json.dumps(report, indent=2))
    _xlsx(report, out / "coverage_report.xlsx")
    for d in documents:
        print(d["document"], {k: d[k] for k in ("jobs", "python_rules", "job_status",
                                                 "orphan_python_rules", "covered_share")})
    print(agent.summary())


if __name__ == "__main__":
    asyncio.run(main())
